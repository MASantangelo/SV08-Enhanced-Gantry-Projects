"""Compare (or sync) a variant's BOM spreadsheet against its assembly STEP file.

Usage (run with FreeCAD's python.exe, which has openpyxl installed):
  python tools/step_bom.py check  "Printed Version"
  python tools/step_bom.py update "Printed Version" [--add-new]
  python tools/step_bom.py tree   "Printed Version"
  python tools/step_bom.py dupes  "Printed Version"

check   Exit 0 if every CAD part maps to a BOM row with the same quantity,
        exit 1 and list the problems otherwise. Never writes anything.
update  Write CAD quantities into the spreadsheet's Qty column (links, notes and
        formatting untouched), then rebuild BOM.md. CAD parts with no BOM row
        are only reported, unless --add-new appends them in a section to sort.
tree    Print the assembly structure, each sub-assembly once. Use it to write rules.
dupes   Check the main STEP and every Base Assemblies STEP for duplicate components:
        Fusion names a copy-pasted (rather than instanced) component "Name (1)", and
        that copy should be another instance of "Name". Exit 1 if any are found.
        check runs this on the main STEP too and fails if it finds any.

How quantities are derived: Fusion's STEP export stores each component's name
(PRODUCT) and every placement of a component inside its parent
(NEXT_ASSEMBLY_USAGE_OCCURRENCE). Walking down from the top assembly and
multiplying counts through sub-assemblies gives the real total of each part.
Rules in tools/bom_rules.json turn raw CAD names into BOM lines.
"""
import argparse
import glob
import json
import os
import re
import subprocess
import sys
from collections import Counter, defaultdict

sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)


# ---------- STEP parsing ----------

def parse_step(path):
    # line breaks carry no meaning in STEP, and long names get wrapped mid-word
    text = open(path, encoding="latin-1").read().replace("\r", "").replace("\n", "")
    data = text[text.index("DATA;") + 5:]
    stmts = {int(m.group(1)): " ".join(m.group(2).split())
             for m in re.finditer(r"#(\d+)\s*=\s*(.*?);\s*(?=#\d+\s*=|ENDSEC)", data, re.S)}

    def refs(s):
        return [int(x) for x in re.findall(r"#(\d+)", s)]

    def strings(s):
        return re.findall(r"'((?:[^']|'')*)'", s)

    product, pdf_product, pd_pdf, edges = {}, {}, {}, []
    for i, s in stmts.items():
        if s.startswith("PRODUCT("):
            st = strings(s)  # (id = Fusion part number, name = component name, ...)
            product[i] = decode(st[1] or st[0])
        elif s.startswith("PRODUCT_DEFINITION_FORMATION"):
            pdf_product[i] = refs(s)[-1]
        elif s.startswith("PRODUCT_DEFINITION("):
            pd_pdf[i] = refs(s)[0]
        elif s.startswith("NEXT_ASSEMBLY_USAGE_OCCURRENCE("):
            r = refs(s)
            edges.append((r[0], r[1]))  # (parent definition, child definition)

    names = {pd: product.get(pdf_product.get(pdf), f"?{pd}") for pd, pdf in pd_pdf.items()}
    kids = defaultdict(Counter)
    has_parent = set()
    for parent, child in edges:
        kids[parent][child] += 1
        has_parent.add(child)
    roots = [pd for pd in pd_pdf if pd not in has_parent]
    return names, kids, roots


def decode(s):
    """Undo STEP string escapes. Sovol's stock part names are GBK bytes in \\X\\hh form."""
    def run(m):
        b = bytes(int(h, 16) for h in re.findall(r"\\X\\([0-9A-F]{2})", m.group(0)))
        try:
            return b.decode("gbk")
        except UnicodeDecodeError:
            return b.decode("latin-1")
    return re.sub(r"(?:\\X\\[0-9A-F]{2})+", run, s.replace("''", "'"))


def clean(name):
    """Merge Fusion copies/mirrors and file-ish suffixes into one part name."""
    n = name.strip()
    while True:
        m = re.sub(r"\s*\((?:\d+|Mirror)\)$", "", n)  # "Name (2)(Mirror)"
        m = re.sub(r"_\d+$", "", m)                    # STEP duplicate "Name_1"
        m = re.sub(r"\.f3d$", "", m, flags=re.I)       # linked "Name.f3d"
        m = re.sub(r"\s+x\d+$", "", m)                 # "Spacer Plate 20mm x2"
        if m == n:
            return n
        n = m


# ---------- rules ----------

def load_rules(variant):
    with open(os.path.join(HERE, "bom_rules.json"), encoding="utf-8") as f:
        all_rules = json.load(f)
    rules = json.loads(json.dumps(all_rules.get("common", {})))
    for key, val in all_rules.get("variants", {}).get(variant, {}).items():
        if isinstance(val, dict):
            rules.setdefault(key, {}).update(val)
        else:
            rules.setdefault(key, []).extend(val)
    return rules


def first_match(patterns, name):
    for pat in patterns:
        if re.search(pat, name):
            return pat
    return None


def cad_quantities(step_path, rules):
    """Return {(category or None, BOM description): qty} plus provenance for reporting."""
    names, kids, roots = parse_step(step_path)
    exclude = rules.get("exclude_assemblies", [])
    units = rules.get("unit_assemblies", {})
    cats = rules.get("category_by_assembly", {})
    rename = rules.get("rename", {})
    rename_patterns = rules.get("rename_patterns", [])
    ignore = rules.get("ignore_parts", [])
    qty, raw = Counter(), defaultdict(set)

    def bom_name(n):
        if n in rename:  # exact renames win over patterns
            return rename[n]
        for pat, repl in rename_patterns:
            if re.search(pat, n):
                return re.sub(pat, repl, n)
        return n

    def walk(pd, mult, cat, only=None):
        n = clean(names[pd])
        if only and not kids[pd] and not re.search(only[1], n):
            cat = only[0]  # a part the category's "only" filter doesn't cover keeps the outer category
        if first_match(exclude, n):
            return
        unit = first_match(units, n)
        if unit:
            desc = units[unit].get("as") or bom_name(n)
            qty[(cat, desc)] += mult
            raw[(cat, desc)].add(n)
            if not units[unit].get("descend", False):
                return
        c = first_match(cats, n)
        if c:
            spec = cats[c]
            if isinstance(spec, dict):  # {"as": category, "only": regex of the part names it applies to}
                only, cat = (cat, spec["only"]), spec["as"]
            else:
                cat = spec
        if kids[pd]:
            for child, k in kids[pd].items():
                walk(child, mult * k, cat, only)
        elif not unit and not first_match(ignore, n):
            desc = bom_name(n)
            qty[(cat, desc)] += mult
            raw[(cat, desc)].add(n)

    for r in roots:  # the root is the top assembly itself (named by part number)
        for child, k in kids[r].items():
            walk(child, k, None)
    return qty, raw


# ---------- spreadsheet ----------

def find_variant_file(variant_dir, pattern, what):
    hits = [p for p in glob.glob(os.path.join(variant_dir, pattern)) if not os.path.basename(p).startswith("~$")]
    if len(hits) != 1:
        sys.exit(f"Expected exactly one {what} ({pattern}) in {variant_dir}, found {len(hits)}: {hits}")
    return hits[0]


def load_bom(xlsx_path):
    import openpyxl
    try:
        wb = openpyxl.load_workbook(xlsx_path)
    except PermissionError:
        sys.exit(f"Can't open {os.path.basename(xlsx_path)}: it's probably open in Excel. Close it and retry.")
    ws = wb.worksheets[0]
    header_row = cols = None
    for row in ws.iter_rows():
        if row[0].value == "Category":
            header_row = row[0].row
            cols = {str(c.value).strip(): c.column for c in row if c.value}
            break
    if not header_row:
        sys.exit("No header row starting with 'Category' found in the spreadsheet.")
    need = ["Category", "Qty Needed", "Description"]
    missing = [c for c in need if c not in cols]
    if missing:
        sys.exit(f"Spreadsheet header is missing columns: {missing}")
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        desc = ws.cell(r, cols["Description"]).value
        if not desc:
            continue
        qty = ws.cell(r, cols["Qty Needed"]).value
        rows.append({"row": r, "cat": str(ws.cell(r, cols["Category"]).value or "").strip(),
                     "desc": str(desc).strip(), "qty": int(qty) if str(qty or "").strip().isdigit() else qty})
    return wb, ws, cols, rows


# ---------- compare ----------

def compare(cad, rows, rules):
    """Match BOM rows to CAD keys. A CAD key with a category only matches a row in that category."""
    bom_only = rules.get("bom_only", [])
    matched_keys, mismatches, bom_missing = set(), [], []
    for row in rows:
        key = (row["cat"], row["desc"])
        if key not in cad:
            key = (None, row["desc"])
        if key in cad and key not in matched_keys:
            matched_keys.add(key)
            row["cad_qty"] = cad[key]
            if row["qty"] != cad[key]:
                mismatches.append(row)
        elif key in cad:
            row["dup"] = True
            bom_missing.append(row)
        elif not first_match(bom_only, row["desc"]):
            bom_missing.append(row)
    cad_missing = {k: q for k, q in cad.items() if k not in matched_keys}
    return mismatches, cad_missing, bom_missing


def report(mismatches, cad_missing, bom_missing, raw):
    ok = not (mismatches or cad_missing or bom_missing)
    if mismatches:
        print(f"\nQUANTITY DIFFERENCES ({len(mismatches)})")
        for r in mismatches:
            print(f"  row {r['row']:>3}  {r['desc']}: BOM {r['qty']} -> CAD {r['cad_qty']}")
    if cad_missing:
        print(f"\nIN CAD BUT NOT IN BOM ({len(cad_missing)})  - add a BOM row, or a rename/ignore rule")
        for (cat, desc), q in sorted(cad_missing.items(), key=lambda x: x[0][1]):
            src = ", ".join(sorted(raw[(cat, desc)] - {desc}))
            where = f" [{cat}]" if cat else ""
            print(f"  {q:>4} x {desc}{where}" + (f"   (CAD name: {src})" if src else ""))
    if bom_missing:
        print(f"\nIN BOM BUT NOT IN CAD ({len(bom_missing)})  - removed from the model, renamed, or add to bom_only")
        for r in bom_missing:
            dup = "  (duplicate row)" if r.get("dup") else ""
            print(f"  row {r['row']:>3}  {r['qty']} x {r['desc']}  [{r['cat']}]{dup}")
    print("\nCHECK PASSED: BOM matches the STEP file." if ok else "\nCHECK FAILED")
    return ok


# ---------- duplicate components ----------

# Stock reference models (Sovol's own parts), reported but not failed on
STOCK_ASSEMBLY = r"^SV08 .* Z Carriage$"


def duplicate_components(step_path):
    """Components named "Name (N)": Fusion's name for a copied component that should have been
    another instance of "Name". Returns [(dup name, instances, parent names, base exists, stock)]."""
    names, kids, roots = parse_step(step_path)
    count, parents = Counter(), defaultdict(set)
    for parent, ck in kids.items():
        for child, k in ck.items():
            count[child] += k
            parents[child].add(parent)
    stock = set()

    def mark(pd):
        for c in kids[pd]:
            if c not in stock:
                stock.add(c)
                mark(c)
    for pd, n in names.items():
        if re.search(STOCK_ASSEMBLY, n.strip()):
            mark(pd)
    all_names = {n.strip() for n in names.values()}
    found = []
    for pd, n in names.items():
        m = re.match(r"^(.*\S)\s*\((\d+)\)$", n.strip())
        if m and pd not in roots:
            par = sorted({names.get(p, "?").strip() for p in parents[pd]})
            found.append((n.strip(), count[pd], par, m.group(1) in all_names, pd in stock))
    return sorted(found, key=lambda f: (f[4], f[0]))


def report_duplicates(step_path):
    found = duplicate_components(step_path)
    own = [f for f in found if not f[4]]
    stock = [f for f in found if f[4]]
    name = os.path.basename(step_path)
    if own:
        print(f"  DUP  {name}: {len(own)} duplicate component(s)")
    else:
        print(f"  OK   {name}: no duplicate components")
    for n, k, par, base, st in own:
        root = re.sub(r"\s*\(\d+\)$", "", n)
        note = (f'copy of "{root}"' if base else
                f'no "{root}" in this file; rename it, or instance the original')
        print(f"         {n}  x{k}  in {', '.join(par)}  ({note})")
    if stock:
        print(f"       ({len(stock)} more inside the stock SV08 Z carriage models, not counted: "
              + ", ".join(sorted({f[0] for f in stock})) + ")")
    return not own


# ---------- tree ----------

def print_tree(step_path):
    names, kids, roots = parse_step(step_path)
    seen = set()

    def show(pd, depth, k):
        n = clean(names[pd])
        pad = "  " * depth
        prefix = f"{k}x " if k > 1 else ""
        if not kids[pd]:
            print(f"{pad}{prefix}{n}")
            return
        if n in seen:
            print(f"{pad}{prefix}[{n}]  (expanded above)")
            return
        seen.add(n)
        print(f"{pad}{prefix}[{n}]")
        for child, ck in sorted(kids[pd].items(), key=lambda x: (bool(kids[x[0]]), clean(names[x[0]]))):
            show(child, depth + 1, ck)

    for r in roots:
        show(r, 0, 1)


# ---------- main ----------

def show_path(p):
    try:
        return os.path.relpath(p, ROOT)
    except ValueError:  # different drive
        return p


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("mode", choices=["check", "update", "tree", "dupes"])
    ap.add_argument("variant", help='variant folder, e.g. "Printed Version"')
    ap.add_argument("--step", help="assembly STEP (default: the single .step in the variant folder)")
    ap.add_argument("--xlsx", help="BOM spreadsheet (default: the single *BOM*.xlsx in the variant folder)")
    ap.add_argument("--add-new", action="store_true", help="update: append CAD parts that have no BOM row")
    a = ap.parse_args()

    variant_dir = os.path.join(ROOT, a.variant)
    step = a.step or find_variant_file(variant_dir, "*.step", "assembly STEP")
    if a.mode == "tree":
        print_tree(step)
        return
    if a.mode == "dupes":
        print("Duplicate components (copies that should be instances):")
        files = [step] + sorted(glob.glob(os.path.join(variant_dir, "Base Assemblies", "*.step")))
        ok = all([report_duplicates(f) for f in files])
        sys.exit(0 if ok else 1)

    rules = load_rules(a.variant)
    xlsx = a.xlsx or find_variant_file(variant_dir, "*BOM*.xlsx", "BOM spreadsheet")
    print(f"STEP: {show_path(step)}\nBOM:  {show_path(xlsx)}")
    cad, raw = cad_quantities(step, rules)
    wb, ws, cols, rows = load_bom(xlsx)
    mismatches, cad_missing, bom_missing = compare(cad, rows, rules)
    ok = report(mismatches, cad_missing, bom_missing, raw)
    if a.mode == "check":
        print("\nDuplicate components (copies that should be instances):")
        if not report_duplicates(step):
            print("\nCHECK FAILED: duplicate components")
            ok = False
        sys.exit(0 if ok else 1)

    changed = 0
    for r in mismatches:
        ws.cell(r["row"], cols["Qty Needed"]).value = r["cad_qty"]
        changed += 1
    if a.add_new and cad_missing:
        r = ws.max_row + 2
        ws.cell(r, 1).value = "New from CAD (sort these)"
        for (cat, desc), q in sorted(cad_missing.items(), key=lambda x: x[0][1]):
            r += 1
            ws.cell(r, cols["Category"]).value = cat or "Unsorted"
            ws.cell(r, cols["Qty Needed"]).value = q
            ws.cell(r, cols["Description"]).value = desc
            changed += 1
    if not changed:
        print("\nNothing to write.")
        return
    try:
        wb.save(xlsx)
    except PermissionError:
        sys.exit(f"Can't save {os.path.basename(xlsx)}: close it in Excel and retry.")
    print(f"\nUPDATED {changed} row(s) in {os.path.basename(xlsx)}")

    if os.path.normcase(os.path.dirname(os.path.abspath(xlsx))) != os.path.normcase(os.path.abspath(variant_dir)):
        print("Spreadsheet is outside the variant folder, so BOM.md was not rebuilt.")
        return
    bom_md = os.path.join(variant_dir, "BOM.md")
    title = a.variant
    if os.path.exists(bom_md):
        first = open(bom_md, encoding="utf-8").readline()
        title = re.sub(r"^#\s*|:\s*Bill of Materials\s*$", "", first.strip()) or title
    subprocess.run([sys.executable, "-I", os.path.join(HERE, "make_bom_md.py"),
                    variant_dir, os.path.basename(xlsx), title], check=True)
    if bom_missing or (cad_missing and not a.add_new):
        print("Quantities synced, but there are unmatched items above. Run check again after fixing them.")


if __name__ == "__main__":
    main()
